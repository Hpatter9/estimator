"""CSV / Excel importers for estimates and room-measurement tables.

Column names are matched loosely, so most exports work as-is. Expected columns (any order):
  estimates: room/section, category, selector, description, quantity, unit, unit price, tax
  rooms:     room, length, width, height, floor sf, ceiling sf, wall sf, perimeter
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from ..models import Estimate, LineItem, Room, Section

ALIASES = {
    "section": ["room", "section", "area", "group", "group description", "location"],
    "category": ["category", "cat", "trade"],
    "selector": ["selector", "sel", "code", "item code"],
    "description": ["description", "desc", "item", "line item"],
    "quantity": ["quantity", "qty"],
    "unit": ["unit", "units", "uom"],
    "unit_price": ["unit price", "unit cost", "price", "rate", "replace"],
    "remove": ["remove"],
    "line_total": ["item amount", "amount", "line total", "total", "extended", "rcv"],
    "tax": ["tax"],
    "length_ft": ["length", "len"],
    "width_ft": ["width"],
    "height_ft": ["height", "ceiling height"],
    "floor_sf": ["floor sf", "floor", "floor area", "area sf", "sf floor"],
    "ceiling_sf": ["ceiling sf", "ceiling", "sf ceiling"],
    "wall_sf": ["wall sf", "walls", "wall area", "sf walls"],
    "perimeter_lf": ["perimeter", "perimeter lf", "lf floor perimeter"],
}


def _norm(h: str) -> str:
    return re.sub(r"[^a-z ]+", " ", str(h or "").lower()).strip()


def _map_headers(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normed = [_norm(h) for h in headers]
    for field, names in ALIASES.items():
        for name in names:
            if name in normed and field not in mapping:
                mapping[field] = normed.index(name)
    return mapping


def _rows(path: Path) -> list[list]:
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [r for r in csv.reader(f) if any(c.strip() for c in r)]
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True, read_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]


def _f(v) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").replace("$", ""))
    except ValueError:
        return None


def read_estimate_sheet(path: Path) -> Estimate:
    rows = _rows(path)
    header_idx = next((i for i, r in enumerate(rows) if "description" in _map_headers(r)), None)
    if header_idx is None:  # e.g. a subcontractor's proposal laid out as a letter, not a line-item table
        return Estimate(title=path.stem, sections=[], source_file=path.name)
    cols = _map_headers(rows[header_idx])
    get = lambda r, k: r[cols[k]] if k in cols and cols[k] < len(r) else None
    sections: dict[str, Section] = {}
    for r in rows[header_idx + 1:]:
        desc = get(r, "description")
        if not desc:
            continue
        name = str(get(r, "section") or "General")
        qty = _f(get(r, "quantity")) or 0.0
        price = (_f(get(r, "unit_price")) or 0.0) + (_f(get(r, "remove")) or 0.0)
        if not price and qty and (total := _f(get(r, "line_total"))):
            price = round(total / qty, 2)  # sheets with only a line total (e.g. "Item Amount")
        sections.setdefault(name, Section(name=name)).items.append(LineItem(
            category=str(get(r, "category") or ""), selector=str(get(r, "selector") or ""),
            description=str(desc), quantity=qty,
            unit=str(get(r, "unit") or "EA"), unit_price=price, tax=_f(get(r, "tax")) or 0.0,
            price_source="history",
        ))
    return Estimate(title=path.stem, sections=list(sections.values()), source_file=path.name)


def read_room_sheet(path: Path) -> list[Room]:
    rows = _rows(path)
    cols = _map_headers(rows[0])
    if "section" not in cols:
        raise ValueError(f"{path.name}: no 'Room' column found")
    rooms = []
    for r in rows[1:]:
        get = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
        if not get("section"):
            continue
        room = Room(name=str(get("section")), **{k: _f(get(k)) for k in
                    ("length_ft", "width_ft", "height_ft", "floor_sf", "ceiling_sf", "wall_sf", "perimeter_lf")})
        rooms.append(fill_derived(room))
    return rooms


def fill_derived(room: Room) -> Room:
    """Compute missing areas from L x W x H when possible."""
    L, W, H = room.length_ft, room.width_ft, room.height_ft
    if L and W:
        room.floor_sf = room.floor_sf or round(L * W, 2)
        room.ceiling_sf = room.ceiling_sf or room.floor_sf
        room.perimeter_lf = room.perimeter_lf or round(2 * (L + W), 2)
    if room.perimeter_lf and H:
        room.wall_sf = room.wall_sf or round(room.perimeter_lf * H, 2)
    return room
