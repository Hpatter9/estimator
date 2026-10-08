from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..models import Estimate

HEADERS = ["Room", "Optional", "Trade", "Category", "Selector", "Description", "Quantity", "Unit", "Unit Price", "Tax", "Total",
           "Price Source", "Note"]


def _rows(est: Estimate):
    for s, i in est.all_items():
        yield [s.name, s.option, i.trade, i.category, i.selector, i.description, i.quantity, i.unit, i.unit_price, i.tax, i.total,
               i.price_source, i.note]


def write_csv(est: Estimate, path: Path) -> Path:
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        w.writerows(_rows(est))
    return path


def write_xlsx(est: Estimate, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Line Items"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    flag = PatternFill("solid", fgColor="FFF2CC")
    for row in _rows(est):
        ws.append(row)
        if row[11] == "ai":  # highlight prices to double-check
            for c in ws[ws.max_row]:
                c.fill = flag
    for col, width in zip("ABCDEFGHIJKLM", [18, 14, 22, 9, 9, 55, 10, 6, 11, 9, 12, 11, 40]):
        ws.column_dimensions[col].width = width
    for r in ws.iter_rows(min_row=2, min_col=9, max_col=11):
        for c in r:
            c.number_format = "#,##0.00"

    summary = wb.create_sheet("Summary")
    for row in [
        ("Estimate", est.title), ("Customer", est.customer), ("Address", est.address),
        ("Claim #", est.claim_number), ("Date", est.estimate_date.isoformat()), (),
        ("Line item total", est.line_item_total), ("Sales tax", est.sales_tax),
        (f"Overhead ({est.overhead_pct:g}%)", est.overhead), (f"Profit ({est.profit_pct:g}%)", est.profit),
        ("Total (base scope)", est.grand_total),
        *[(f"Optional: {o}", est.scope_total(o)) for o in est.options],
    ]:
        summary.append(list(row))
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["B"].width = 40
    wb.save(path)
    return path
